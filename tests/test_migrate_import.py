"""Schema migrations (an old database brought up to date, twice) and the spreadsheet importer with malformed and
duplicate input. The workbook here is built in the test from made-up numbers in the real sheet's layout."""
from __future__ import annotations
import sqlite3
from datetime import date, datetime
import pytest
from app import db
from app.importer import import_workbook, seed_favorites, verify
from app.migrate import COLUMNS, migrate

# The schema as it was on 2026-09-17, before any migration existed, plus an early `recurring` without anchor_day.
OLD = db.SCHEMA + """
CREATE TABLE recurring(id INTEGER PRIMARY KEY, label TEXT NOT NULL, category_id INTEGER NOT NULL,
  from_account_id INTEGER, to_account_id INTEGER, amount INTEGER NOT NULL, what TEXT NOT NULL DEFAULT '',
  freq TEXT NOT NULL, next_date TEXT NOT NULL, horizon_days INTEGER NOT NULL DEFAULT 45,
  active INTEGER NOT NULL DEFAULT 1, created_at TEXT NOT NULL DEFAULT (datetime('now')));
"""


def schema(con):
    return {r[0]: r[1] for r in con.execute("SELECT name, sql FROM sqlite_master ORDER BY name")}


def old_db(path):
    con = sqlite3.connect(path)
    con.executescript(OLD)
    con.execute("INSERT INTO accounts(name,kind,start_balance) VALUES('Checking','cash',10000)")
    con.execute("INSERT INTO categories(name,type) VALUES('Roth IRA','Saving'),('Other Income','Money in'),"
                "('Food','Spending')")
    con.execute("INSERT INTO transactions(date,what,category_id,from_account_id,amount) VALUES('2026-09-02','x',3,1,500)")
    con.execute("INSERT INTO recurring(label,category_id,from_account_id,amount,freq,next_date)"
                " VALUES('R',3,1,100,'monthly','2026-10-31')")
    con.execute("INSERT INTO settings VALUES('start_month','2026-08-01')")
    con.commit()
    con.close()


def test_old_schema_migrates_and_keeps_its_data(tmp_path):
    path = str(tmp_path / "old.db")
    old_db(path)
    con = db.connect(path)
    for table, cols in COLUMNS.items():
        have = {r[1] for r in con.execute(f"PRAGMA table_info({table})")}
        assert {c for c, _ in cols} <= have, table
    row = dict(con.execute("SELECT * FROM transactions").fetchone())
    assert row["amount"] == 500 and row["note"] == "" and row["tags"] == "" and row["client_id"] is None
    assert con.execute("SELECT anchor_day FROM recurring").fetchone()[0] is None
    settings = dict(con.execute("SELECT key, value FROM settings").fetchall())
    assert settings["roth_category"] == "1" and settings["interest_category"] == "2"  # pinned by id
    assert settings["start_month"] == "2026-08-01" and settings["ef_months"] == "6"  # defaults don't overwrite
    assert {"budgets", "saved_views", "tx_client"} <= set(schema(con))


def test_migration_twice_changes_nothing(tmp_path):
    path = str(tmp_path / "twice.db")
    old_db(path)
    con = db.connect(path)
    before = schema(con)
    migrate(con)
    migrate(con)
    con.commit()
    assert schema(con) == before
    assert con.execute("SELECT COUNT(*) FROM settings WHERE key='roth_category'").fetchone()[0] == 1


def test_a_renamed_pin_is_not_repinned(tmp_path):
    """Once pinned, renaming Roth IRA (or making a new category with that name) must not move the pin."""
    path = str(tmp_path / "pin.db")
    old_db(path)
    con = db.connect(path)
    con.execute("UPDATE categories SET name='Retirement' WHERE id=1")
    con.execute("INSERT INTO categories(name,type) VALUES('Roth IRA','Saving')")
    migrate(con)
    assert db.setting(con, "roth_category") == "1"


def test_recurring_from_the_old_table_still_posts(tmp_path):
    from app.recurring import generate
    path = str(tmp_path / "rec.db")
    old_db(path)
    con = db.connect(path)
    assert generate(con, date(2026, 12, 1)) == 3
    # Without anchor_day (templates made before it existed) the day follows the last date, so the 31st drifts to
    # the 30th after a short month. Templates saved since then carry the anchor.
    assert [r[0] for r in con.execute("SELECT date FROM transactions WHERE recurring_id=1 ORDER BY date")] == \
        ["2026-10-31", "2026-11-30", "2026-12-30"]


def test_client_id_is_unique(tmp_path):
    con = db.connect(str(tmp_path / "u.db"))
    con.execute("INSERT INTO accounts(name,kind) VALUES('A','cash')")
    con.execute("INSERT INTO categories(name,type) VALUES('C','Spending')")
    ins = "INSERT INTO transactions(date,category_id,from_account_id,amount,client_id) VALUES('2026-09-01',1,1,5,?)"
    con.execute(ins, ("abc12345",))
    con.execute(ins, (None,))
    con.execute(ins, (None,))
    with pytest.raises(sqlite3.IntegrityError):
        con.execute(ins, ("abc12345",))


# ---------- importer ----------

openpyxl = pytest.importorskip("openpyxl")
ACCOUNTS = ["Chase Checking", "SoFi Checking", "Amex HYSA", "SoFi HYSA", "Chase CC", "Amex CC", "HSA"]
CATS = [("Paycheck", "Money in"), ("Groceries", "Spending"), ("Roth IRA", "Saving"), ("HYSA Transfer", "Transfer"),
        ("Student Loan", "Loan")]


def workbook(path, rows):
    """A made-up workbook in the Money + Log layout the importer reads."""
    wb = openpyxl.Workbook()
    log = wb.active
    log.title = "Log"
    money = wb.create_sheet("Money")
    for i, name in enumerate(ACCOUNTS):
        log[f"M{5 + i}"] = name
    for i, (name, typ) in enumerate(CATS):
        log[f"J{5 + i}"], log[f"K{5 + i}"] = name, typ
    for i, r in enumerate(rows):
        for col, v in zip("ABCDEF", r, strict=True):
            log[f"{col}{5 + i}"] = v
    for i, (name, bal) in enumerate([("Chase Checking", 1000.5), ("Amex HYSA", 5000), ("Chase CC", 120.1)]):
        money[f"A{123 + i}"], money[f"C{123 + i}"] = name, bal
    money["C110"], money["D110"], money["C111"], money["D111"] = 8000, 0.05, 2000, 0.03
    money["AJ4"] = datetime(2026, 9, 1)
    money["A27"], money["A28"] = datetime(2026, 8, 1), datetime(2026, 9, 1)
    money["AA27"], money["AB27"], money["AC27"] = 1500.25, 3000, 700
    money["C103"], money["C100"] = 4, 7000
    wb.save(path)


ROWS = [
    (datetime(2026, 8, 3), "Pay", "Paycheck", None, "Chase Checking", 2000),
    (datetime(2026, 8, 4), " Store ", "Groceries", "Chase CC", None, 12.345),  # third decimal: half-up
    (None, "blank date is skipped", "Groceries", "Chase CC", None, 1),
    (date(2026, 9, 1), "Roth", "Roth IRA", "Chase Checking", None, 500),
    (datetime(2026, 9, 2), "To savings", "HYSA Transfer", "Chase Checking", "Amex HYSA", 250),
]


def test_import_reads_the_layout(tmp_path):
    path = tmp_path / "book.xlsx"
    workbook(path, ROWS)
    con = db.connect(str(tmp_path / "i.db"))
    out = import_workbook(con, str(path))
    assert out["transactions"] == 4 and out["categories"] == 5
    kinds = dict(con.execute("SELECT name, kind FROM accounts").fetchall())
    assert kinds["Chase CC"] == "card" and kinds["HSA"] == "investment" and kinds["Amex HYSA"] == "cash"
    assert kinds["Roth IRA"] == "investment" and kinds["Student loans"] == "loan"
    starts = dict(con.execute("SELECT name, start_balance FROM accounts").fetchall())
    assert starts["Chase Checking"] == 100050 and starts["Chase CC"] == 12010 and starts["Student loans"] == 1000000
    rate = con.execute("SELECT loan_rate FROM accounts WHERE kind='loan'").fetchone()[0]
    assert rate == pytest.approx((8000 * 0.05 + 2000 * 0.03) / 10000)
    amounts = [r[0] for r in con.execute("SELECT amount FROM transactions ORDER BY date")]
    assert amounts == [200000, 1235, 50000, 25000]
    assert con.execute("SELECT what FROM transactions WHERE amount=1235").fetchone()[0] == "Store"
    typed = con.execute("SELECT a.name, t.month, t.balance FROM typed_balances t JOIN accounts a ON a.id=t.account_id")
    assert {r[0]: (r[1], r[2]) for r in typed} == {"Roth IRA": ("2026-08-01", 150025), "401(k)": ("2026-08-01", 300000),
                                                  "HSA": ("2026-08-01", 70000)}
    assert db.setting(con, "ef_months") == "4" and db.setting(con, "roth_limit") == "700000"
    assert seed_favorites(con) == 1


def test_verify_catches_a_one_cent_difference(tmp_path):
    path = tmp_path / "book.xlsx"
    workbook(path, ROWS)
    con = db.connect(str(tmp_path / "v.db"))
    import_workbook(con, str(path))
    # write the engine's own figures into the sheet, then nudge one by a cent
    from app.engine import month_range, month_table
    rows = month_table(db.load_accounts(con), db.load_categories(con), db.load_txns(con), db.load_typed(con),
                       month_range(date(2026, 8, 1), date(2026, 9, 1)))
    wb = openpyxl.load_workbook(path)
    money = wb["Money"]
    for r, row in zip((27, 28), rows, strict=True):
        money[f"B{r}"], money[f"Q{r}"], money[f"U{r}"] = row.money_in / 100, row.spent / 100, row.left_over / 100
    money["Q27"] = rows[0].spent / 100 + 0.01
    wb.save(path)
    bad = verify(con, str(path))
    assert any(b.startswith("Q27") for b in bad) and not any(b.startswith("B27") for b in bad)


def test_import_refuses_a_database_with_entries(tmp_path):
    path = tmp_path / "book.xlsx"
    workbook(path, ROWS)
    con = db.connect(str(tmp_path / "d.db"))
    import_workbook(con, str(path))
    with pytest.raises(SystemExit):
        import_workbook(con, str(path))
    assert con.execute("SELECT COUNT(*) FROM transactions").fetchone()[0] == 4


def test_malformed_files_fail_before_writing(tmp_path):
    con = db.connect(str(tmp_path / "m.db"))
    junk = tmp_path / "junk.xlsx"
    junk.write_bytes(b"not a spreadsheet")
    from zipfile import BadZipFile
    from openpyxl.utils.exceptions import InvalidFileException
    with pytest.raises((BadZipFile, InvalidFileException)):
        import_workbook(con, str(junk))
    wb = openpyxl.Workbook()
    wb.active.title = "Money"
    wb.save(tmp_path / "nolog.xlsx")
    with pytest.raises(KeyError):
        import_workbook(con, str(tmp_path / "nolog.xlsx"))
    assert con.execute("SELECT COUNT(*) FROM accounts").fetchone()[0] == 0


def test_an_unknown_category_stops_the_import_uncommitted(tmp_path):
    path = tmp_path / "book.xlsx"
    workbook(path, ROWS + [(datetime(2026, 9, 3), "?", "Not a category", "Chase CC", None, 5)])
    dbp = str(tmp_path / "c.db")
    con = db.connect(dbp)
    with pytest.raises(KeyError):
        import_workbook(con, str(path))
    con.close()
    assert sqlite3.connect(dbp).execute("SELECT COUNT(*) FROM transactions").fetchone()[0] == 0
